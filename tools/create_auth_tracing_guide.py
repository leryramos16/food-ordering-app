from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Food_Ordering_Authentication_Code_Tracing_Guide.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_ORANGE = "FFF3E8"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_index])
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for index, text in enumerate(headers):
        table.rows[0].cells[index].text = text
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Code"]
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run(text)
    return paragraph


def add_steps(doc, steps):
    for number, (title, body) in enumerate(steps, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(f"{number}. {title}  ")
        run.bold = True
        paragraph.add_run(body)


def add_file_ref(doc, path, line, purpose):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"{path}:{line}")
    run.bold = True
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    paragraph.add_run(f"  {purpose}")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = BLACK
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12

styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = BLACK
styles["Title"].paragraph_format.space_after = Pt(12)
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size, before, after in (
    ("Heading 1", 19, 12, 8),
    ("Heading 2", 13.5, 10, 5),
    ("Heading 3", 11.5, 8, 4),
):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = BLACK
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

code_style = styles.add_style("Code", 1)
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.5)
code_style.font.color.rgb = BLACK
code_style.paragraph_format.left_indent = Inches(0.22)
code_style.paragraph_format.right_indent = Inches(0.12)
code_style.paragraph_format.space_before = Pt(4)
code_style.paragraph_format.space_after = Pt(8)
code_style.paragraph_format.line_spacing = 1.0
code_style._element.get_or_add_pPr().append(OxmlElement("w:shd"))
code_style._element.pPr[-1].set(qn("w:fill"), LIGHT_GRAY)

doc.add_paragraph("Food Ordering Authentication Code Tracing Guide", style="Title")
subtitle = doc.add_paragraph("Flutter client to Laravel Sanctum API")
subtitle.runs[0].font.size = Pt(15)
subtitle.runs[0].font.bold = True
subtitle.paragraph_format.space_after = Pt(20)

doc.add_heading("Purpose", level=1)
doc.add_paragraph(
    "This guide explains how authentication travels through the food ordering app, from a tap in Flutter to validation and token creation in Laravel, then back to secure storage on the phone. Read it beside the source code and follow the file references in order."
)
doc.add_paragraph(
    "The central idea is simple: Flutter owns presentation and local session state; Laravel owns identity, validation, passwords, authorization, and database records. The client never chooses the authenticated user ID. Laravel identifies the user from the Sanctum bearer token."
)

doc.add_heading("System at a glance", level=1)
add_table(doc, ["Layer", "Responsibility", "Main code"], [
    ("Flutter presentation", "Collects input and shows progress or errors", "login_screen.dart, register_screen.dart"),
    ("Flutter state", "Tracks current user, loading state, and errors", "auth_controller.dart"),
    ("Flutter data", "Builds requests and saves or clears the token", "auth_service.dart, api_client.dart"),
    ("Laravel routes", "Maps URLs and applies Sanctum middleware", "routes/api.php"),
    ("Laravel application", "Validates, authenticates, creates users and tokens", "Requests, AuthController, UserResource"),
    ("Database", "Stores users, hashed passwords, and token hashes", "users, personal_access_tokens"),
], [1.25, 3.0, 2.2])

doc.add_heading("One complete round trip", level=2)
add_code(doc, "Tap button -> AuthController -> AuthService -> ApiClient -> HTTP request\n-> Laravel route -> FormRequest -> AuthController -> User model / database\n-> JSON response -> ApiClient -> AuthService -> secure token storage\n-> AuthController notifies UI -> protected HomeScreen")

doc.add_page_break()
doc.add_heading("Project structure and dependency direction", level=1)
doc.add_paragraph("Dependencies point inward. A screen calls state, state calls a service, and the service calls the reusable API and storage classes. Laravel follows a similar separation: route, request validation, controller, model or service, then resource response.")
add_code(doc, "mobile/lib/\n  app.dart\n  core/\n    api_config.dart\n    api_client.dart\n    api_exception.dart\n    token_storage.dart\n  features/auth/\n    domain/app_user.dart\n    data/auth_service.dart\n    state/auth_controller.dart\n    presentation/login_screen.dart\n    presentation/register_screen.dart\n  features/home/presentation/home_screen.dart\n\napp/Http/\n  Controllers/Api/AuthController.php\n  Requests/LoginRequest.php\n  Requests/RegisterRequest.php\n  Resources/UserResource.php\nroutes/api.php\napp/Models/User.php")

doc.add_heading("Application startup trace", level=2)
add_steps(doc, [
    ("Flutter starts", "main.dart creates FoodOrderingEstApp."),
    ("Dependencies are assembled", "app.dart constructs TokenStorage, ApiClient, AuthService, and AuthController once."),
    ("Session restoration starts", "AuthController.restoreSession reads the stored token and calls GET /auth/me when a token exists."),
    ("The UI chooses a screen", "A loading spinner appears while checking. A valid user opens HomeScreen; otherwise LoginScreen opens."),
])
add_file_ref(doc, "mobile/lib/app.dart", 20, "Dependency construction and the authentication gate.")
add_file_ref(doc, "mobile/lib/features/auth/state/auth_controller.dart", 17, "Startup session restoration.")

doc.add_heading("Why this separation matters", level=2)
doc.add_paragraph("Screens do not know HTTP headers or JSON shapes. Laravel controllers do not know Flutter widgets. Token storage is isolated so it can later be replaced or mocked without rewriting login and signup screens.")

doc.add_page_break()
doc.add_heading("Signup trace", level=1)
doc.add_paragraph("Signup creates a customer and immediately signs that customer in. The server, not Flutter, assigns the customer role and active state.")
add_steps(doc, [
    ("User submits the form", "RegisterScreen validates name, email, password length, and password confirmation before making a request."),
    ("State enters loading mode", "AuthController.register clears the old error, disables the button, and calls AuthService.register."),
    ("The service creates JSON", "AuthService trims text and sends name, email, optional phone, password, password_confirmation, and device_name."),
    ("ApiClient sends HTTP", "POST /api/auth/register is sent with Accept and Content-Type set to application/json."),
    ("Laravel selects the route", "routes/api.php maps the public endpoint to AuthController.register."),
    ("Laravel validates input", "RegisterRequest enforces unique email and phone, confirmed password, and a minimum of eight characters."),
    ("Laravel creates the user", "The controller forces role=customer and is_active=true. The User model's hashed cast hashes the password before storage."),
    ("Sanctum creates a token", "createToken returns the plain token once; Laravel stores only its hash in personal_access_tokens."),
    ("Flutter saves the session", "AuthService stores the token with flutter_secure_storage and converts user JSON into AppUser."),
    ("The protected screen opens", "AuthController sets user and notifies listeners. app.dart rebuilds and selects HomeScreen."),
])

doc.add_heading("Signup request", level=2)
add_code(doc, "POST /api/auth/register\nAccept: application/json\nContent-Type: application/json\n\n{\n  \"name\": \"Johnlery Ramos\",\n  \"email\": \"customer@example.com\",\n  \"phone\": \"09161234567\",\n  \"password\": \"example-password\",\n  \"password_confirmation\": \"example-password\",\n  \"device_name\": \"food-ordering-mobile\"\n}")
add_file_ref(doc, "mobile/lib/features/auth/presentation/register_screen.dart", 130, "Form submission entry point.")
add_file_ref(doc, "app/Http/Requests/RegisterRequest.php", 15, "Authoritative server validation.")
add_file_ref(doc, "app/Http/Controllers/Api/AuthController.php", 17, "User and token creation.")

doc.add_page_break()
doc.add_heading("Login trace", level=1)
doc.add_paragraph("Login proves identity using email and password. A failure returns validation-style JSON without revealing whether the email or password was wrong.")
add_steps(doc, [
    ("User taps Sign in", "LoginScreen runs local checks and calls AuthController.login."),
    ("Flutter posts credentials", "AuthService sends the email, password, and device name to /auth/login."),
    ("Laravel validates shape", "LoginRequest requires a valid email string and a password string."),
    ("Laravel finds the account", "AuthController normalizes the email to lowercase and queries users."),
    ("Laravel verifies the password", "Hash::check compares the submitted password to the stored password hash."),
    ("Laravel checks account status", "An inactive account is rejected even if the password is correct."),
    ("A device token is issued", "authenticatedResponse creates a new Sanctum token and returns safe user fields."),
    ("Flutter persists the token", "The service saves the token, creates AppUser, and the UI switches to HomeScreen."),
])

doc.add_heading("Successful authentication response", level=2)
add_code(doc, "{\n  \"success\": true,\n  \"message\": \"Logged in successfully.\",\n  \"data\": {\n    \"user\": {\"id\": 1, \"name\": \"...\", \"email\": \"...\", \"phone\": null, \"role\": \"customer\"},\n    \"token\": \"1|plain-token-returned-once\",\n    \"token_type\": \"Bearer\"\n  }\n}")
doc.add_paragraph("UserResource controls which user fields leave the API. Password, password hash, active flag, and internal token records are not returned.")
add_file_ref(doc, "mobile/lib/features/auth/presentation/login_screen.dart", 102, "Login form submission.")
add_file_ref(doc, "app/Http/Controllers/Api/AuthController.php", 38, "Credential and active-account checks.")
add_file_ref(doc, "app/Http/Resources/UserResource.php", 10, "Safe public user shape.")

doc.add_page_break()
doc.add_heading("Token use and session restoration", level=1)
doc.add_paragraph("The token is the mobile app's proof of authentication after login. It replaces sending the password again. Treat it like a temporary key: send it only to the API over HTTPS in production and remove it at logout.")

doc.add_heading("Authenticated request trace", level=2)
add_steps(doc, [
    ("A service requests authentication", "It calls ApiClient with authenticated: true."),
    ("ApiClient reads secure storage", "TokenStorage returns the saved token."),
    ("ApiClient attaches the header", "Authorization: Bearer TOKEN is added to the request."),
    ("Sanctum middleware runs", "auth:sanctum hashes or resolves the supplied token and finds its user."),
    ("Laravel exposes the user", "$request->user() becomes the authenticated User model."),
])
add_code(doc, "GET /api/auth/me\nAccept: application/json\nAuthorization: Bearer 1|plain-token-returned-once")

doc.add_heading("App restart trace", level=2)
add_steps(doc, [
    ("The app starts", "app.dart calls restoreSession immediately."),
    ("No token", "AuthService returns null and LoginScreen appears."),
    ("Token exists", "AuthService calls GET /auth/me."),
    ("Token is valid", "The API returns the current user and HomeScreen appears without another login."),
    ("Token is invalid or unreachable", "The current implementation clears the local token and returns to LoginScreen."),
])
doc.add_paragraph("Study note: clearing the token on every restore error keeps this first version simple. A later production version may distinguish an invalid token from a temporary offline connection so users are not signed out during a network outage.")
add_file_ref(doc, "mobile/lib/core/api_client.dart", 41, "Bearer header creation.")
add_file_ref(doc, "mobile/lib/core/token_storage.dart", 3, "Secure token read, save, and clear operations.")
add_file_ref(doc, "routes/api.php", 12, "Sanctum-protected route group.")

doc.add_page_break()
doc.add_heading("Logout trace", level=1)
doc.add_paragraph("Logout removes the current device token on both sides. Other devices remain logged in because only currentAccessToken is deleted.")
add_steps(doc, [
    ("User taps the logout icon", "HomeScreen calls AuthController.logout."),
    ("Flutter calls the API", "AuthService posts to /auth/logout with the bearer token."),
    ("Sanctum authenticates", "The middleware resolves the token and user before the controller runs."),
    ("Laravel revokes one token", "currentAccessToken()->delete() removes the token record used by this request."),
    ("Flutter always clears local storage", "A finally block clears the token even when the network request fails."),
    ("State becomes unauthenticated", "AuthController sets user to null, notifies listeners, and LoginScreen replaces HomeScreen."),
])

doc.add_heading("Why finally is used", level=2)
doc.add_paragraph("The local app should still sign out if the server is temporarily unavailable. The tradeoff is that a failed server request may leave an unreachable server token active until it expires or is revoked later.")
add_file_ref(doc, "mobile/lib/features/auth/data/auth_service.dart", 59, "Server logout followed by guaranteed local clearing.")
add_file_ref(doc, "app/Http/Controllers/Api/AuthController.php", 70, "Current token deletion.")

doc.add_heading("Multiple devices", level=2)
add_table(doc, ["Action", "Phone A", "Phone B", "Database"], [
    ("Login on both", "Token A stored", "Token B stored", "Two token rows"),
    ("Logout on Phone A", "Token A cleared", "Still authenticated", "Token A row deleted"),
    ("Revoke all tokens later", "Signed out", "Signed out", "All user token rows deleted"),
], [1.35, 1.6, 1.6, 1.9])

doc.add_page_break()
doc.add_heading("Authenticated order ownership", level=1)
doc.add_paragraph("The important security improvement is that order creation no longer accepts user_id from Flutter. A malicious client could change a supplied user_id. The controller now takes the user from the verified Sanctum token.")
add_code(doc, "Flutter sends: restaurant_id, address_id, payment_method, notes, items\nFlutter does not send: user_id\n\nLaravel route: auth:sanctum -> OrderController.store\nLaravel identity: $user = $request->user()\nOrderService: confirms address.user_id == authenticated user.id")

doc.add_heading("Order request trace", level=2)
add_steps(doc, [
    ("Bearer token identifies the customer", "Sanctum resolves the user before OrderController runs."),
    ("StoreOrderRequest validates order shape", "Restaurant, address, payment method, and item values are checked."),
    ("OrderController passes the trusted user", "The controller passes $request->user() and validated order data into OrderService."),
    ("OrderService verifies ownership", "The address must match both address_id and authenticated user_id."),
    ("Server calculates trusted totals", "Menu prices come from the database, preventing client-side price manipulation."),
    ("Database transaction commits", "Order and item snapshots are saved together or rolled back together."),
])
add_file_ref(doc, "routes/api.php", 33, "Order endpoint protected by Sanctum.")
add_file_ref(doc, "app/Http/Controllers/Api/OrderController.php", 14, "Authenticated user passed to the service.")
add_file_ref(doc, "app/Services/OrderService.php", 17, "Ownership, availability, prices, totals, and transaction.")

doc.add_heading("Trust boundary", level=2)
add_table(doc, ["Flutter may suggest", "Laravel must decide"], [
    ("Email and password input", "Whether credentials are valid"),
    ("Address ID", "Whether the address belongs to this customer"),
    ("Menu item IDs and quantities", "Availability, restaurant match, and current prices"),
    ("Displayed subtotal", "Authoritative subtotal, fees, discounts, and total"),
], [3.15, 3.3])

doc.add_page_break()
doc.add_heading("Validation and error flow", level=1)
doc.add_paragraph("Laravel returns structured errors. ApiClient converts the first validation message into ApiException, AuthController stores its text, and the screen displays it beneath the fields.")
add_code(doc, "Laravel 422 response\n{\n  \"message\": \"The email has already been taken.\",\n  \"errors\": {\"email\": [\"The email has already been taken.\"]}\n}\n\nApiException.fromResponse -> AuthController.errorMessage -> red text in screen")

doc.add_heading("Common responses", level=2)
add_table(doc, ["Status", "Meaning", "Typical cause"], [
    ("200", "Request succeeded", "Login, profile, or logout completed"),
    ("201", "Resource created", "Signup created a new customer"),
    ("401", "Unauthenticated", "Missing, invalid, or revoked bearer token"),
    ("422", "Validation failed", "Duplicate email, weak password, or bad credentials"),
    ("500", "Server failure", "Database configuration or an unhandled backend error"),
], [0.75, 2.0, 3.7])

doc.add_heading("The two connection errors encountered", level=2)
add_table(doc, ["Message", "Layer reached", "Meaning and fix"], [
    ("SocketException connection refused 127.0.0.1:8000", "Flutter network", "The phone could not reach Laravel. Start php artisan serve and restore adb reverse tcp:8000 tcp:8000."),
    ("SQLSTATE 1045 access denied", "Laravel database", "Flutter reached Laravel, but Laravel's MySQL credentials were rejected. Local development was switched to the existing SQLite database and migrations were run."),
], [2.2, 1.25, 3.0])
add_file_ref(doc, "mobile/lib/core/api_exception.dart", 1, "API error extraction and display message.")

doc.add_page_break()
doc.add_heading("Security decisions in this version", level=1)
add_table(doc, ["Decision", "Why it matters"], [
    ("Passwords use the User model hashed cast", "Plain passwords are never stored in the database."),
    ("Generic incorrect-credentials message", "Does not reveal which email addresses are registered."),
    ("Role forced to customer during signup", "A caller cannot register itself as an owner or administrator."),
    ("Active status checked at login", "Disabled accounts cannot obtain new tokens."),
    ("Sanctum protects private routes", "Only requests with valid bearer tokens reach their controllers."),
    ("Token stored with flutter_secure_storage", "More appropriate than plain shared preferences for credentials."),
    ("UserResource limits response fields", "Sensitive or internal model fields do not leak to Flutter."),
    ("Order identity comes from token", "Users cannot create orders under another user ID."),
], [2.55, 3.9])

doc.add_heading("Production follow ups", level=2)
for item in (
    "Use HTTPS and remove Android cleartext traffic permission for production builds.",
    "Restore MySQL with a dedicated account and valid least-privilege credentials before deployment.",
    "Add email verification, password reset, token expiration policy, and rate limiting when the app grows.",
    "Distinguish offline startup from invalid-token startup if offline usability becomes important.",
    "Never commit .env, tokens, real passwords, or production database credentials.",
):
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("What is intentionally not included yet", level=2)
doc.add_paragraph("There is no social login, refresh-token subsystem, role management UI, password reset screen, or complex architecture framework. The current design is deliberately small enough to study while still supporting real separate users and per-device sessions.")

doc.add_page_break()
doc.add_heading("Recommended study order", level=1)
add_steps(doc, [
    ("Start with the routes", "Read routes/api.php to see the public and protected entry points."),
    ("Read Laravel validation", "Compare LoginRequest and RegisterRequest with their Flutter forms."),
    ("Trace AuthController", "Follow register, login, me, logout, then authenticatedResponse."),
    ("Inspect User and UserResource", "Compare database-capable fields with the smaller API response."),
    ("Move to Flutter ApiClient", "See headers, JSON conversion, bearer tokens, and errors."),
    ("Read AuthService", "See endpoint-specific request bodies and token lifecycle."),
    ("Read AuthController state", "Follow loading, user, error, and notifyListeners."),
    ("Finish with screens and app.dart", "Observe how widgets react to state without owning backend logic."),
    ("Run the tests", "Use php artisan test and flutter test, then deliberately try duplicate email and wrong password cases."),
])

doc.add_heading("Trace checklist", level=2)
for item in (
    "I can identify where local form validation ends and server validation begins.",
    "I can explain why password_confirmation is sent only during signup.",
    "I can locate where the password is hashed.",
    "I can explain why the plain token is returned only once.",
    "I can find where Authorization Bearer is added.",
    "I can explain how app.dart chooses LoginScreen or HomeScreen.",
    "I can show why orders cannot impersonate another user.",
    "I can distinguish network, validation, authentication, and database errors.",
):
    doc.add_paragraph("[ ] " + item)

doc.add_heading("Verification commands", level=2)
add_code(doc, "# Backend\nphp artisan test\nphp artisan route:list --path=api\n\n# Flutter\ncd mobile\nflutter analyze\nflutter test\n\n# Connected Android phone using USB forwarding\nadb reverse tcp:8000 tcp:8000\nflutter run -d DEVICE_ID --dart-define=API_BASE_URL=http://127.0.0.1:8000/api")

doc.add_heading("Tests already covering the flow", level=2)
doc.add_paragraph("tests/Feature/AuthTest.php verifies signup and token creation, login and profile retrieval, logout token deletion, bad credential rejection, and authentication on order creation. mobile/test/app_user_test.dart verifies conversion of the Laravel user JSON into the Flutter AppUser model.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Food Ordering Authentication Code Tracing Guide")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
